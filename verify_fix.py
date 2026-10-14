import openpyxl

print("--- Checking Kotak Bank (a84f3374_Kotak_Bank_OD_NoR.xlsx) ---")
wb_k = openpyxl.load_workbook("output/a84f3374_Kotak_Bank_OD_NoR.xlsx", data_only=True)
ws_k = wb_k["Transactions"]
print("Kotak Row 2:", [c.value for c in ws_k[2]])
missing = sum(1 for r in range(2, ws_k.max_row + 1) if ws_k.cell(r, 2).value is None or ws_k.cell(r, 3).value is None)
print(f"Kotak missing dates across {ws_k.max_row - 1} rows: {missing}")
ws_kv = wb_k["Validation"]
for r in range(1, 10):
    k, v = ws_kv.cell(r, 1).value, ws_kv.cell(r, 2).value
    if k in ("Balance-chain pass rate", "Date-field completeness", "Rows flagged CHECK", "Statement completeness"):
        print(f"  {k}: {v}")

print("\n--- Checking Bank of Baroda (cc2a6f28_banking_bank_of_baroda.xlsx) ---")
wb_b = openpyxl.load_workbook("output/cc2a6f28_banking_bank_of_baroda.xlsx", data_only=True)
ws_b = wb_b["Transactions"]
print("Baroda S.No 354 (Row 355):", [c.value for c in ws_b[355]])
print("Baroda S.No 988 (Row 989):", [c.value for c in ws_b[989]])
missing_narrs = sum(1 for r in range(2, ws_b.max_row + 1) if not ws_b.cell(r, 4).value or str(ws_b.cell(r, 4).value).strip() == "")
print(f"Baroda blank narrations across {ws_b.max_row - 1} rows: {missing_narrs}")
ws_bv = wb_b["Validation"]
for r in range(1, 10):
    k, v = ws_bv.cell(r, 1).value, ws_bv.cell(r, 2).value
    if k in ("Balance-chain pass rate", "Date-field completeness", "Rows flagged CHECK", "Statement completeness"):
        print(f"  {k}: {v}")

